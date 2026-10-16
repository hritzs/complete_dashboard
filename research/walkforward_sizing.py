"""Annual past-only maximum-profit sizing; conditional on frozen legacy entries."""
import hashlib,json,pickle
from pathlib import Path
import numpy as np
import pandas as pd
from sizing_drawdown_risk import risk_fit, stress_history
from wf_sizing_features import DAILY,OUT,ROOT

CORE=['Table','IV_Ratio','Str_Ratio','Build_IV','DTE','Norm_OG','Adj_IV_Chg','Fixed_IDV']
NUMERIC=['IV','RV10','RV20','VRP10','VRP20','SPREAD20','RV_TREND','IV_Ratio','Str_Ratio',
         'Norm_OG','Adj_IV_Chg','Fixed_IDV','u_range_bps','str_range_pct','str_chg_abs','spread_pct','S1_score','S2_score']
CW_ONLY=set(NUMERIC[7:])


def rank(reference,values):
    ref=np.sort(np.asarray(reference,float));v=np.asarray(values,float)
    l=np.searchsorted(ref,v,'left');r=np.searchsorted(ref,v,'right')
    return np.where(r>l,(l+r+1)/2,l)/len(ref)


def measure(p):
    p=np.asarray(p,float);eq=np.cumsum(p);dd=np.maximum.accumulate(np.r_[0.,eq])[1:]-eq
    return float(p.sum()),float(dd.max()),float(p.min())


def project(raw,lower,upper):
    """Exact piecewise-linear bounded projection; no future data in apply()."""
    raw=np.asarray(raw,float);lo=np.asarray(lower,float);hi=np.asarray(upper,float);target=len(raw)
    if (lo>hi).any() or lo.sum()>target+1e-10 or hi.sum()<target-1e-10:raise ValueError('Infeasible mean-one bounds')
    if abs(lo.sum()-target)<1e-12:offset=float(np.min(lo-raw))
    elif abs(hi.sum()-target)<1e-12:offset=float(np.max(hi-raw))
    else:
        breaks=np.r_[lo-raw,hi-raw];events=np.r_[np.ones(target),-np.ones(target)]
        order=np.argsort(breaks,kind='stable');breaks=breaks[order];slope=np.cumsum(events[order])
        totals=lo.sum()+np.r_[0.,np.cumsum(np.diff(breaks)*slope[:-1])]
        j=int(np.searchsorted(totals,target,'left'))
        j=min(max(j,1),len(breaks)-1);i=j-1
        offset=float(breaks[j] if abs(totals[j]-target)<1e-12 else breaks[i]+(target-totals[i])/slope[i])
    size=np.clip(raw+offset,lo,hi)
    if abs(size.mean()-1)>1e-10:raise AssertionError('Projection budget error')
    return size,offset


def scalefit(a,train):
    a=np.asarray(a,float)
    med=np.array([np.median(a[train,j][np.isfinite(a[train,j])]) if np.isfinite(a[train,j]).any() else 0. for j in range(a.shape[1])])
    filled=np.where(np.isfinite(a),a,med);mean=filled[train].mean(0);std=filled[train].std(0);std[std<1e-12]=1.
    return np.clip((filled-mean)/std,-3,3),{'Median':med.tolist(),'Mean':mean.tolist(),'Std':std.tolist()}


def prepare():
    files=sorted((ROOT/'output/backtest_relax/run_027bf0249b/days').glob('*.pkl'))
    d=pd.DataFrame([pickle.load(f.open('rb')) for f in files]).sort_values('Date').reset_index(drop=True)
    d.Net_PnL=d.Net_PnL.fillna(0.);d.Leg=d.Leg.fillna('')
    f=pd.read_csv(OUT/'strict_features.csv',dtype={'Source_mtime_ns':str})
    meta=json.loads((OUT/'strict_features.json').read_text())
    assert meta['CSV_sha256']==hashlib.sha256((OUT/'strict_features.csv').read_bytes()).hexdigest()
    assert meta['Code_sha256']==hashlib.sha256((ROOT/'wf_sizing_features.py').read_bytes()).hexdigest()
    assert set(f.Date)==set(d.loc[d.Leg.isin(['CW','NW']),'Date']) and not f.Date.duplicated().any()
    valid=f.Valid.fillna(False).astype(bool)
    assert (pd.to_datetime(f.loc[valid,'Max_feature_tick'])<pd.to_datetime(f.loc[valid,'Decision_time'])).all()
    daily=pd.read_csv(DAILY);daily.Date=pd.to_datetime(daily.Date).dt.strftime('%Y-%m-%d')
    daily=daily.drop_duplicates('Date',keep='last').sort_values('Date').reset_index(drop=True)
    daily['PrevDate']=pd.to_datetime(daily.PrevDate);daily['Prev2Date']=pd.to_datetime(daily.Prev2Date)
    returns=np.log(daily.Future/daily.Future_Prev2).where(lambda x:np.isfinite(x)&x.abs().lt(.2))
    for n in (5,10,20):daily[f'RV{n}']=returns.rolling(n,min_periods=n).std()*np.sqrt(252)
    refs=daily.set_index('Date');traded=d.Leg.isin(['CW','NW']);td=d[traded].reset_index(drop=True)
    # Replace all intraday market predictors; retain instrument/entry/outcome metadata.
    td=td.drop(columns=[c for c in ('IV_Ratio','Str_Ratio','Build_IV','Norm_OG','Adj_IV_Chg','Fixed_IDV','DTE_idx') if c in td])
    td=td.merge(f.drop(columns=['Leg','Entry_Time']),on='Date',validate='one_to_one')
    refrows=[]
    for row in td.itertuples():
        r=refs.loc[row.Date];day=pd.Timestamp(row.Date)
        if not (r.PrevDate<day and r.Prev2Date<day):raise ValueError('Non-past daily reference '+row.Date)
        pure,wprev=float(r.IDV_Pure),float(r.IDV_With_Prev)
        if not(np.isfinite(pure) and np.isfinite(wprev)):
            old=daily[(daily.Date<row.Date)&daily.IDV_Pure.notna()&daily.IDV_With_Prev.notna()&daily.PrevDate.ge(day-pd.Timedelta(days=30))&daily.PrevDate.lt(day)]
            if len(old):pure,wprev=map(float,old.sort_values('PrevDate').iloc[-1][['IDV_Pure','IDV_With_Prev']])
        idv=.2*pure+.5*wprev+.3*r.Adj_IV
        gap=(r.PrevDate-r.Prev2Date).days>1;base=r.Prev2_Adj_IV if gap else r.Prev2_IV
        change=r.Final_IV_Chg if 'Final_IV_Chg' in daily else r.Adj_IV_Chg if gap else r.IV_Chg
        raw=float(row.Raw_DTE) if pd.notna(row.Raw_DTE) else np.nan
        if np.isfinite(raw):
            calendar=int(np.ceil(raw));business=int(np.busday_count(np.datetime64(row.Date,'D'),np.datetime64(row.Date,'D')+np.timedelta64(calendar,'D')))
            tdte=max(0.,raw-1.) if calendar>business else raw
            biv=row.Raw_IV*np.sqrt(raw/tdte) if tdte>0 else np.nan
            di=int(np.clip(np.ceil(tdte)-1,0,6))
        else:biv=np.nan;di=np.nan
        refrows.append({'Build_IV':biv,'IV':biv,'IV_Ratio':biv/idv if idv>0 else np.nan,
            'Str_Ratio':r.Straddle/row.Straddle if pd.notna(row.Straddle) and row.Straddle>0 else np.nan,
            'Fixed_IDV':idv,'Adj_IV_Chg':change*.1/base if base!=0 else np.nan,
            'Norm_OG':((row.Opening_underlying-r.Future)/r.Future*100)*.19/r.Adj_IV if r.Adj_IV>0 else np.nan,
            'DTE_idx':di,'DTE':min(di+1,5) if pd.notna(di) else np.nan,
            'Table':{'09:16':0,'09:17':1,'09:18':2}.get(row.Stage,3),
            'RV10':r.RV10,'RV20':r.RV20,'RV_TREND':r.RV5/r.RV20 if r.RV20>0 else np.nan,
            'Reference_prev_date':str(r.PrevDate.date()),'Reference_prev2_date':str(r.Prev2Date.date())})
    td=pd.concat([td,pd.DataFrame(refrows)],axis=1)
    for n in (10,20):td[f'VRP{n}']=td.IV/td[f'RV{n}']
    td['SPREAD20']=td.IV-td.RV20;td['Year']=pd.to_datetime(td.Date).dt.year
    dd,streak=stress_history(d.Net_PnL.to_numpy(float));td['Prior_flat_DD']=dd[traded];td['Prior_flat_loss_streak']=streak[traded]
    return d,td,files


def scores(td,train):
    cw=td.Leg.eq('CW').to_numpy();fit=train&cw
    a=td[CORE].to_numpy(float);med=np.array([np.nanmedian(a[fit,j]) if np.isfinite(a[fit,j]).any() else 0. for j in range(len(CORE))])
    a=np.where(np.isfinite(a),a,med);p=td.Net_PnL.to_numpy(float);mu=p[fit].mean();s1=np.zeros(len(td));buckets=[]
    for j,name in enumerate(CORE):
        edges=np.unique(a[fit,j]) if name in ('Table','DTE') else np.quantile(a[fit,j],[.2,.4,.6,.8])
        bins=a[:,j] if name in ('Table','DTE') else np.searchsorted(edges,a[:,j],side='right')
        levels=edges if name in ('Table','DTE') else range(5)
        effects=[]
        for level in levels:
            mask=fit&(bins==level);n=int(mask.sum());value=(p[mask].mean()-mu)*n/(n+20) if n else 0.
            s1[cw&(bins==level)]+=value;effects.append([float(level),float(value),n])
        buckets.append({'Feature':name,'Edges':edges.tolist(),'Effects':effects})
    # Table is represented by three indicators in S2; other core fields numeric.
    numeric=[j for j,name in enumerate(CORE) if name!='Table']
    b=np.column_stack([a[:,numeric]]+[(a[:,0]==k).astype(float) for k in (1,2,3)])
    mean=b[fit].mean(0);std=b[fit].std(0);std[std<1e-12]=1.
    z=(b-mean)/std;beta=np.linalg.solve(z[fit].T@z[fit]+10*np.eye(z.shape[1]),z[fit].T@(p[fit]-mu))
    s2=z@beta;s1[~cw]=np.nan;s2[~cw]=np.nan
    return s1,s2,{'Core':CORE,'Median':med.tolist(),'Buckets':buckets,'S2_mean':mean.tolist(),'S2_std':std.tolist(),'S2_beta':beta.tolist(),'PnL_mean':float(mu)}


def base_design(td,train):
    s1,s2,score_model=scores(td,train);data=td.copy();data['S1_score']=s1;data['S2_score']=s2
    cw=data.Leg.eq('CW').to_numpy();a=data[NUMERIC].to_numpy(float)
    for j,name in enumerate(NUMERIC):
        if name in CW_ONLY:a[:,j]=np.where(cw,a[:,j],np.nan)
    x,transform=scalefit(a,train);names=list(NUMERIC)
    for col,levels in [('Stage',['09:17','09:18','09:19','09:20+']),('DTE',[2,3,4,5]),('Weekday',[1,2,3,4,5,6])]:
        vals=pd.to_datetime(data.Date).dt.weekday.to_numpy() if col=='Weekday' else data[col].to_numpy()
        for level in levels:
            x=np.column_stack([x,(cw&(vals==level)).astype(float)]);names.append(f'CW {col}={level}')
    x=np.column_stack([x,(cw&data.Event.fillna(False).to_numpy(bool)).astype(float),(~cw).astype(float)]);names+=['CW Event','NW']
    return x,names,{'Numeric_features':NUMERIC,**transform,'Profit_scores':score_model}


def risk_design(td,train,x,names):
    stress=np.column_stack([np.clip(td.Prior_flat_DD.to_numpy()/75,0,3),np.clip(td.Prior_flat_loss_streak.to_numpy()/3,0,3)])
    x=np.column_stack([x,stress]);names=names+['Prev_flat_DD','Prev_flat_loss_streak']
    keep=np.array([n not in ('S1_score','S2_score') for n in names]);risknames=list(np.array(names)[keep]);a=x[:,keep]
    risk,coef=risk_fit(a[train],td.loc[train,'Net_PnL'].to_numpy(),risknames,td[train].reset_index(drop=True))
    beta=coef.Logistic_beta.to_numpy();gamma=coef.Downside_beta.to_numpy()
    loss=1/(1+np.exp(-np.clip(beta[0]+a@beta[1:],-25,25)));down=np.maximum(gamma[0]+a@gamma[1:],0)
    refs={};parts=[];cw=td.Leg.eq('CW').to_numpy();missing=np.zeros(len(td),bool)
    for col in ('u_range_bps','str_range_pct','spread_pct'):
        v=td[col].to_numpy(float);known=cw&np.isfinite(v);ref=v[train&known]
        if not len(ref):raise ValueError('Missing training instability sample')
        refs[col]=np.sort(ref).tolist();parts.append(np.where(known,rank(ref,np.nan_to_num(v)),.5));missing|=cw&~known
    lr=rank(loss[train],loss);dr=rank(down[train],down);instability=np.mean(parts,axis=0)
    score=.4*lr+.4*dr+.2*instability;threshold=float(risk.Risk_threshold.iloc[0]);dicey=(score>=threshold)|missing
    np.testing.assert_allclose(score[train],risk.Risk_score,atol=1e-10,rtol=0)
    readings=pd.DataFrame({'Date':td.Date,'Loss_probability':loss,'Expected_downside':down,'Loss_rank':lr,'Downside_rank':dr,
        'Instability':instability,'Risk_score':score,'Risk_threshold':threshold,'Dicey':dicey,'Missing_CW':missing})
    extra,transform=scalefit(np.column_stack([loss,down]),train)
    model={'Features':risknames,'Logistic_beta':beta.tolist(),'Downside_beta':gamma.tolist(),
        'Loss_reference':np.sort(loss[train]).tolist(),'Downside_reference':np.sort(down[train]).tolist(),
        'Instability_references':refs,'Threshold':threshold,'Allocation_risk_transform':transform}
    return np.column_stack([x,extra]),names+['Risk_Loss_probability','Risk_Expected_downside'],readings,model


def optimize(x,p,lower,upper,risk_policy):
    seed=73 if risk_policy else 42;rng=np.random.default_rng(seed);flat=measure(p);n=x.shape[1]
    feasible=lower.sum()<=len(p)+1e-10 and upper.sum()>=len(p)-1e-10
    default=np.zeros(n)
    fallback={'Beta':default.tolist(),'Offset':0.,'Status':'Risk-only fallback' if risk_policy else 'Flat fallback',
              'Feasible_mean_one':bool(feasible),'Qualified':False,'Evaluations':0}
    if not feasible:return fallback
    seen=None;seenvalue=-np.inf;evaluations=0
    def assess(beta):
        nonlocal seen,seenvalue,evaluations
        evaluations+=1;size,offset=project(1+np.einsum('ij,j->i',x,beta),lower,upper);net,dd,worst=measure(p*size)
        qualified=(not risk_policy) or (net>=flat[0]-1e-8 and dd<=flat[1]+1e-8 and worst>=flat[2]-1e-8)
        reg=1e-5*np.mean(beta**2)
        if qualified:value=net/flat[0]-reg
        else:value=-1e6-max(0,1-net/flat[0])-max(0,dd/flat[1]-1)-max(0,(flat[2]-worst)/abs(flat[2]))-reg
        r={'beta':beta.copy(),'offset':offset,'value':value,'Qualified':qualified,'Train_net':net,'Train_MDD':dd,'Train_worst':worst}
        if qualified and value>seenvalue:seen=r;seenvalue=value
        return r
    candidates=[assess(default)]
    for scale in ((.04,.12,.25) if risk_policy else (.025,.075,.15,.3)):
        for _ in range(20 if risk_policy else 24):candidates.append(assess(np.clip(rng.normal(0,scale,n),-1,1)))
    for j in range(n):
        for direction in (-1,1):
            b=default.copy();b[j]=direction*.08;candidates.append(assess(b))
    for current in sorted(candidates,key=lambda r:r['value'],reverse=True)[:2]:
        for step in ((.12,.06,.03,.015,.0075) if risk_policy else (.12,.06,.03,.015,.0075,.00375)):
            for _ in range(3):
                changed=False
                for j in range(n):
                    best=current
                    for direction in (-1,1):
                        b=current['beta'].copy();b[j]=np.clip(b[j]+direction*step,-1,1);trial=assess(b)
                        if trial['value']>best['value']+1e-10:best=trial
                    if best is not current:current=best;changed=True
                if not changed:break
    if seen is None:fallback['Evaluations']=evaluations;return fallback
    return {'Beta':seen['beta'].tolist(),'Offset':seen['offset'],'Status':'Optimized on prior years','Feasible_mean_one':True,
        'Qualified':True,'Evaluations':evaluations,**{k:seen[k] for k in ('Train_net','Train_MDD','Train_worst')}}


def bounds(td,dicey=None):
    lower=np.full(len(td),.75);upper=np.where(td.Leg.eq('CW')&td.Stage.eq('09:16'),1.,1.15)
    if dicey is not None:upper=np.where(dicey,np.minimum(upper,.85),upper)
    return lower,upper


def fold(td,year):
    train=td.Year.lt(year).to_numpy();apply=td.Year.eq(year).to_numpy()
    if train.sum()<50:raise ValueError('Insufficient training')
    x,names,base_model=base_design(td,train);xr,rnames,risk,risk_model=risk_design(td,train,x,names)
    lo,hi=bounds(td);rl,rh=bounds(td,risk.Dicey.to_numpy())
    models={};sizes={};features={}
    for label,xx,nn,ll,hh,isrisk in [('WF10',x,names,lo,hi,False),('WF11',xr,rnames,rl,rh,True)]:
        fitted=optimize(xx[train],td.loc[train,'Net_PnL'].to_numpy(),ll[train],hh[train],isrisk)
        value=np.clip(1+xx@np.asarray(fitted['Beta'])+fitted['Offset'],ll,hh)
        models[label]={'Features':nn,**fitted};sizes[label]=value
        features[label]=pd.DataFrame({'Date':td.Date,**{n:xx[:,j] for j,n in enumerate(nn)}})
    model={'Year':year,'Train_until':f'{year-1}-12-31','Last_training_entry':str(td.loc[train,'Date'].max()),
        'First_scored_entry':str(td.loc[apply,'Date'].min()),'Training_entries':int(train.sum()),
        'Scored_entries':int(apply.sum()),'Base_model':base_model,'Risk_model':risk_model,'Policies':models}
    assert model['Last_training_entry']<model['First_scored_entry']
    return apply,sizes,risk,features,model


def main():
    OUT.mkdir(parents=True,exist_ok=True);d,td,files=prepare();traded=d.Leg.isin(['CW','NW']).to_numpy()
    sizes={'Flat size 1':np.ones(len(td)),'WF10 maximum profit':np.ones(len(td)),'WF11 dicey0.85 maximum profit':np.ones(len(td))}
    models=[];risks=[];input_rows=[];coefs=[];cutoffs=[]
    for year in sorted(td.Year.unique()):
        if year==2021:continue
        print(f'Fitting prior-years-only fold{year}',flush=True)
        apply,pred,risk,features,model=fold(td,int(year));models.append(model)
        (OUT/f'model_{year}.json').write_text(json.dumps(model,indent=2)+'\n')
        for label,name in [('WF10','WF10 maximum profit'),('WF11','WF11 dicey0.85 maximum profit')]:
            sizes[name][apply]=pred[label][apply];p=model['Policies'][label]
            cutoffs.append({'Year':int(year),'Policy':name,'Train_until':model['Train_until'],'Last_training_entry':model['Last_training_entry'],
                'First_scored_entry':model['First_scored_entry'],'Training_entries':model['Training_entries'],
                'Scored_entries':model['Scored_entries'],'Status':p['Status'],'Feasible_mean_one':p['Feasible_mean_one'],
                'Qualified_training':p['Qualified'],'Frozen_offset':p['Offset'],'Evaluations':p['Evaluations']})
            coefs.extend({'Year':int(year),'Candidate':name,'Feature':n,'Coefficient':b,'Frozen_offset':p['Offset']} for n,b in zip(p['Features'],p['Beta']))
            f=features[label].loc[apply].copy();f['Policy']=name;f['Year']=int(year);input_rows.append(f)
        r=risk.loc[apply].copy();r['Year']=int(year);r['Training_cutoff']=model['Train_until'];risks.append(r)
        print({k:(model['Policies'][k]['Status'],float(pred[k][apply].mean())) for k in pred},flush=True)
    risk=pd.concat(risks,ignore_index=True);entries=td.merge(risk.drop(columns='Year'),on='Date',how='left',validate='one_to_one')
    # Keep warmup readings explicitly unestimated.
    entries['Warmup']=entries.Year.eq(2021);entries['Dicey']=entries.Dicey.fillna(False).astype(bool)
    for name,s in sizes.items():entries[name+'_size']=s;entries[name+'_PnL']=td.Net_PnL.to_numpy()*s
    source_daily=d.merge(entries[['Date','Loss_probability','Expected_downside','Loss_rank','Downside_rank','Instability','Risk_score','Risk_threshold','Dicey','Warmup']],on='Date',how='left',validate='one_to_one')
    fullsizes={name:np.ones(len(d)) for name in sizes}
    for name,s in sizes.items():fullsizes[name][traded]=s
    daily_sizes=pd.DataFrame({'Date':d.Date,**fullsizes});daily_pnl=pd.DataFrame({'Date':d.Date,**{name:d.Net_PnL.to_numpy()*s for name,s in fullsizes.items()}})
    scored_mask=d.Date.ge('2022-01-01').to_numpy();paths={'Date':d.loc[scored_mask,'Date'].to_numpy()}
    for name in fullsizes:
        equity=np.cumsum(daily_pnl.loc[scored_mask,name].to_numpy(float))
        paths[name+'_Equity']=equity;paths[name+'_DD']=np.maximum.accumulate(np.r_[0.,equity])[1:]-equity
    summary=[];yearly=[];quarterly=[]
    for name,s in fullsizes.items():
        for scope,mask in [('All_history',np.ones(len(d),bool)),('Warmup_2021',d.Date.lt('2022-01-01').to_numpy()),('Scored_2022plus',d.Date.ge('2022-01-01').to_numpy())]:
            p=(d.Net_PnL.to_numpy()*s)[mask];net,dd,worst=measure(p);mean=float(s[mask&traded].mean());flat=measure(d.Net_PnL.to_numpy()[mask])
            summary.append({'Candidate':name,'Period':scope,'Net':net,'MDD':dd,'Worst_day':worst,'Net_per_MDD':net/dd,
                'Mean_size':mean,'Cuts':int((s[mask&traded]<1-1e-8).sum()),'Boosts':int((s[mask&traded]>1+1e-8).sum()),
                'Entries':int((mask&traded).sum()),'Same_mean_flat_net':flat[0]*mean,'Same_mean_flat_MDD':flat[1]*mean,
                'Extra_vs_same_mean_flat':net-flat[0]*mean,
                'Qualifies':bool(name!='Flat size 1' and net>flat[0] and net/dd>flat[0]/flat[1] and worst>=flat[2])})
        for column,periods,target in [('Year',pd.to_datetime(d.Date).dt.year,yearly),('Quarter',pd.to_datetime(d.Date).dt.to_period('Q').astype(str),quarterly)]:
            for per in sorted(set(periods)):
                mask=np.asarray(periods==per);net,dd,worst=measure((d.Net_PnL.to_numpy()*s)[mask])
                target.append({'Candidate':name,column:per,'Net':net,'MDD':dd,'Worst_day':worst,'Mean_size':float(s[mask&traded].mean())})
    risk_coefs=[];references=[];transforms=[];s1_rows=[];s2_rows=[];core_medians=[]
    for model in models:
        year=model['Year'];rm=model['Risk_model'];bm=model['Base_model'];sm=bm['Profit_scores']
        risk_coefs.extend({'Year':year,'Feature':name,'Logistic_beta':b,'Downside_beta':g}
            for name,b,g in zip(['Intercept',*rm['Features']],rm['Logistic_beta'],rm['Downside_beta']))
        for variable,values in {'Loss_probability':rm['Loss_reference'],'Expected_downside':rm['Downside_reference'],**rm['Instability_references']}.items():
            references.extend({'Year':year,'Variable':variable,'Sorted_index':i+1,'Value':v,'Exact_value_text':format(v,'.17g'),
                'Risk_threshold':rm['Threshold']} for i,v in enumerate(values))
        for family,names,scaler in [('Base',bm['Numeric_features'],bm),('Risk_allocation',['Loss_probability','Expected_downside'],rm['Allocation_risk_transform'])]:
            transforms.extend({'Year':year,'Family':family,'Feature':n,'Median':med,'Mean':mu,'Std':std,'Clip':3}
                for n,med,mu,std in zip(names,scaler['Median'],scaler['Mean'],scaler['Std']))
        core_medians.extend({'Year':year,'Feature':n,'Imputation_median':med} for n,med in zip(sm['Core'],sm['Median']))
        for bucket in sm['Buckets']:
            s1_rows.extend({'Year':year,'Feature':bucket['Feature'],'Edges':json.dumps(bucket['Edges']),
                'Level':level,'Shrunk_effect':effect,'Training_count':count,'Shrink_count':20,'CW_training_PnL_mean':sm['PnL_mean']}
                for level,effect,count in bucket['Effects'])
        s2_names=[n for n in sm['Core'] if n!='Table']+[f'Table={k}' for k in (1,2,3)]
        s2_rows.extend({'Year':year,'Feature':n,'Mean':mu,'Std':std,'Coefficient':b,'Lambda':10,'CW_training_PnL_mean':sm['PnL_mean']}
            for n,mu,std,b in zip(s2_names,sm['S2_mean'],sm['S2_std'],sm['S2_beta']))
    formulas=[('Loss model','y=1[net PnL<0]; minimize sum(log(1+exp(z))-y*z)+20/2*sum(beta_j^2); intercept unpenalized'),
        ('Loss probability','p=1/(1+exp(-clip(beta0+X*beta,-25,25)))'),
        ('Downside target','d=max(-net PnL,0); winners contribute zero; expected downside is not VaR or worst loss'),
        ('Downside model','gamma=(Xc.T*Xc+50I)^-1*Xc.T*(d-mean(d)); gamma0=mean(d)-mean(X)*gamma; prediction=max(0,gamma0+X*gamma)'),
        ('Reference rank','N=prior training reference count; L=count(ref<value), E=count(ref=value); rank=(L+(E+1)/2)/N if E>0 else L/N'),
        ('Instability','I=(rank(underlying range)+rank(straddle LTP range)+rank(spread))/3; missing or NW component=.5'),
        ('Risk','R=.4*rank(loss probability)+.4*rank(expected downside)+.2*I; threshold=training R 75th percentile'),
        ('Dicey','R>=saved threshold OR missing CW instability inputs; WF11 hard upper=.85; no WF10 dicey upper'),
        ('Sizing inputs','Median-fill using training median; (value-training mean)/training std, clipped[-3,3]; fixed indicators'),
        ('Raw allocation','raw=1+sum(X_j*b_j); b chosen on prior training PnL only; allocation b differs from risk beta/gamma'),
        ('Final allocation','size=clip(raw+frozen training offset,.75,upper); upper1.15,09:16 CW1,WF11 dicey.85'),
        ('Fallback','WF11 infeasible/no qualifying training fit: .85 dicey else1; WF10 flat fallback1; actual future mean reported'),
        ('Net and DD','daily net=original saved net*size; equity=cumulative net; DD=max(0,all prior equity peaks)-equity; MDD=max(DD)'),
        ('Tax','Original LUT hurdles13/5/2/2/6 unchanged; tax is a cell selection hurdle, not a per-trade PnL subtraction'),
        ('Precision','Saved JSON models and CSV inputs retain float precision; reference text preserves 17-digit values. New sizing Excel formulas replay standardized saved inputs; new risk fits are replayed independently in Python.')]
    tables={'Candidates':pd.DataFrame(summary),'By_year':pd.DataFrame(yearly),'By_quarter':pd.DataFrame(quarterly),
        'Daily_risk':source_daily,'Trade_comparison':entries,'Daily_sizes':daily_sizes,'Daily_PnL':daily_pnl,
        'Training_cutoffs':pd.DataFrame(cutoffs),'Coefficients':pd.DataFrame(coefs),'Model_inputs':pd.concat(input_rows,ignore_index=True),
        'Scored_equity_DD':pd.DataFrame(paths),'Risk_Coefficients':pd.DataFrame(risk_coefs),'Risk_References':pd.DataFrame(references),
        'Feature_Transforms':pd.DataFrame(transforms),'Score_Imputation':pd.DataFrame(core_medians),
        'S1_Buckets':pd.DataFrame(s1_rows),'S2_Coefficients':pd.DataFrame(s2_rows),'Formula_Guide':pd.DataFrame(formulas,columns=['Part','Formula'])}
    notes=['Past-only sizing-overlay rerun of10/11; strict prefix intraday features; prior daily reference dates enforced.',
        '2021 warmup size1; each following calendar year fits completed prior years only; no full-history weights/warm starts.',
        'Training mean-one offset frozen in scoring year; actual mean reported, no future normalization.',
        'Legacy base LUTs/entries used wider history; saved fills/costs and dated reference fields are not independently certified causal.',
        'New forward overlay fixes fitting/normalization leakage, not end-to-end entry selection bias; previously researched years are not untouched holdouts.',
        'All original entry taxes preserved; archived09:16 CW cap1; WF11 mandatory dicey<=.85; low.75/global1.15.',
        'WF10 has no mandatory dicey restriction; its risk columns are comparison diagnostics only. WF11 uses the risk flag as a hard cap.',
        'Prefix OTM IV/DTE, straddle ratios and LTP ranges replace cached intraday predictors, so this is a repaired-input rerun, not identical feature values.',
        'If WF11 training feasibility/qualification fails, report explicit risk-only fallback; do not relax risk caps.',
        'Size multipliers scale existing event weights and saved net costs/PnL linearly; no whole-lot or quantity-specific execution replay.']
    tables={'Read_me':pd.DataFrame({'Method':notes}),**tables}
    def finish_excel(writer):
        from openpyxl.chart import LineChart,Reference
        for ws in writer.book.worksheets:
            ws.freeze_panes='B2';ws.auto_filter.ref=ws.dimensions
            ws.column_dimensions['A'].width=24
        chartws=writer.book.create_sheet('Charts');source=writer.book['Scored_equity_DD']
        for metric,anchor in [('Equity','A1'),('DD','A30')]:
            chart=LineChart();chart.title='2022+ '+metric+' (conditional sizing overlay)';chart.width=28;chart.height=14
            for j,cell in enumerate(source[1],1):
                if str(cell.value).endswith('_'+metric):chart.add_data(Reference(source,min_col=j,max_col=j,min_row=1,max_row=source.max_row),titles_from_data=True)
            chart.set_categories(Reference(source,min_col=1,min_row=2,max_row=source.max_row));chartws.add_chart(chart,anchor)
    with pd.ExcelWriter(OUT/'walkforward_comparison.xlsx',engine='openpyxl') as w:
        for sheet,data in tables.items():data.to_excel(w,sheet_name=sheet,index=False)
        finish_excel(w)
    for label,name in [('wf10','WF10 maximum profit'),('wf11','WF11 dicey0.85 maximum profit')]:
        subset=['Flat size 1',name]
        with pd.ExcelWriter(OUT/(label+'.xlsx'),engine='openpyxl') as w:
            for sheet,data in tables.items():
                if sheet=='Daily_sizes':data=daily_sizes[['Date',*subset]]
                elif sheet=='Daily_PnL':data=daily_pnl[['Date',*subset]]
                elif sheet in ('Candidates','By_year','By_quarter','Coefficients'):data=data[data.Candidate.isin(subset)]
                elif sheet=='Training_cutoffs':data=data[data.Policy.eq(name)]
                elif sheet=='Model_inputs':data=data[data.Policy.eq(name)]
                elif sheet=='Scored_equity_DD':data=data[['Date',*[n+'_'+metric for n in subset for metric in ('Equity','DD')]]]
                elif label=='wf10' and sheet=='Daily_risk':sheet='Daily';data=d.merge(td,on='Date',how='left',suffixes=('_base',''))
                data.to_excel(w,sheet_name=sheet,index=False)
            finish_excel(w)
    for name,data in [('summary',tables['Candidates']),('trade_comparison',entries),('training_cutoffs',tables['Training_cutoffs']),('daily_sizes',daily_sizes),('daily_pnl',daily_pnl),('prepared_inputs',td)]:data.to_csv(OUT/(name+'.csv'),index=False)
    scored=tables['Candidates'][tables['Candidates'].Period.eq('Scored_2022plus')]
    print(scored.to_string(index=False),flush=True)
    table=['| Policy | Net PnL | Max DD | Worst day | Net/DD | Mean size | Cuts | Boosts |',
        '|---|---:|---:|---:|---:|---:|---:|---:|']
    table += [f'| {r.Candidate} | {r.Net:,.2f} | {r.MDD:,.2f} | {r.Worst_day:,.2f} | {r.Net_per_MDD:.2f} | {r.Mean_size:.6f} | {r.Cuts} | {r.Boosts} |' for r in scored.itertuples()]
    text=['# Walk-forward sizing rerun','',*['- '+note for note in notes],'','Scored 2022+ (617 entries; 2021 warmup excluded):','',*table,'',
        'Compare forward mean/ex-post same-mean flat diagnostics; these are not a live normalization rule.',
        f'Input coverage: {int(td.Valid.fillna(False).sum())}/{len(td)} valid strict prefixes.']
    (OUT/'REPORT.md').write_text('\n'.join(text)+'\n')
    inputs=[Path(__file__),ROOT/'wf_sizing_features.py',ROOT/'WALKFORWARD_SIZING_SPEC.md',ROOT/'sizing_drawdown_risk.py',
        ROOT/'sizing_normalized_optimization.py',ROOT/'tick_entry_features.py',ROOT/'test_walkforward_sizing.py',ROOT/'WALKFORWARD_EXPLAINED.md',
        DAILY,OUT/'strict_features.csv',OUT/'strict_features.json',*files]
    manifest={'Inputs':{str(p.relative_to(ROOT)) if p.is_relative_to(ROOT) else str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in inputs},
        'Outputs':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in OUT.iterdir() if p.is_file() and p.name!='manifest.json' and not p.name.startswith('pilot')}}
    (OUT/'manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')


if __name__=='__main__':main()
