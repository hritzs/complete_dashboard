"""Audit stored Excel periods, daily paths and spreadsheet-engine risk calculations."""
import hashlib
import json
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'output/research_ledger'


def frame(book,name):
    it=book[name].iter_rows(values_only=True); columns=next(it)
    return pd.DataFrame(it,columns=columns)


def main():
    manifest=json.loads((OUT/'manifest.json').read_text())
    for p,digest in manifest['Input_sha256'].items():
        path=Path(p) if Path(p).is_absolute() else ROOT/p
        assert hashlib.sha256(path.read_bytes()).hexdigest()==digest,p
    for p,digest in manifest['Output_sha256'].items():
        if p in ('Risk_Formulas_Calculated.xlsx','verification.json') and manifest.get('Stage')!='verified': continue
        assert hashlib.sha256((OUT/p).read_bytes()).hexdigest()==digest,p
    book=load_workbook(OUT/'GammaShort_Research_Ledger.xlsx',read_only=True,data_only=True)
    registry=json.loads((ROOT/'research_iterations.json').read_text()); count=0; periods=0
    for config in registry:
        prefix=config['id']; print('Verifying iteration',prefix,flush=True)
        daily=frame(book,prefix+'_Daily'); trades=frame(book,prefix+'_Trades'); summary=frame(book,prefix+'_Summary')
        catalog=frame(book,prefix+'_Policies'); dates=pd.to_datetime(daily.Date); td=pd.to_datetime(trades.Date)
        counts={}
        for row in catalog.to_dict('records'):
            name=row['Policy']; p=daily[name+'_PnL'].to_numpy(float)
            running=peak=0.; mdd=0.; equity=[]; drawdowns=[]
            for i,value in enumerate(p):
                running+=value; peak=max(peak,running); draw=peak-running; mdd=max(mdd,draw)
                equity.append(running); drawdowns.append(draw)
            np.testing.assert_allclose(equity,daily[name+'_Equity'].to_numpy(float),atol=1e-8)
            np.testing.assert_allclose(drawdowns,daily[name+'_DD'].to_numpy(float),atol=1e-8)
            ref=summary[summary.Policy.eq(name)&summary.Period.eq('All_history')].iloc[0]
            np.testing.assert_allclose([sum(p),mdd],[ref.Net,ref.MDD],atol=1e-8)
            np.testing.assert_allclose(trades[name+'_PnL'].sum(),sum(p),atol=1e-8)
            executed=trades[name+'_Action'].ne('skip')
            assert int(executed.sum())==ref.Entries
            counts[name]=executed
            count+=1
        for sheet in ('Quarterly','Yearly','Monthly','Financial_year'):
            data=frame(book,prefix+'_'+sheet)
            for row in data.to_dict('records'):
                period=str(row['Period'])
                if sheet=='Quarterly':
                    year,quarter=map(int,period.split('Q')); start=pd.Timestamp(year,3*quarter-2,1); end=start+pd.DateOffset(months=3)
                elif sheet=='Yearly': start=pd.Timestamp(int(period),1,1); end=start+pd.DateOffset(years=1)
                elif sheet=='Monthly': start=pd.Timestamp(period+'-01'); end=start+pd.DateOffset(months=1)
                else: start=pd.Timestamp(int(period[2:6]),4,1); end=start+pd.DateOffset(years=1)
                mask=dates.ge(start)&dates.lt(end); tmask=td.ge(start)&td.lt(end)
                p=daily.loc[mask,row['Policy']+'_PnL'].to_numpy(float)
                peak=running=mdd=0.
                for value in p:
                    running+=value; peak=max(peak,running); mdd=max(mdd,peak-running)
                np.testing.assert_allclose([sum(p),mdd],[row['Net'],row['MDD']],atol=1e-8)
                assert int((counts[row['Policy']]&tmask).sum())==row['Entries']
                periods+=1
    # Verify the separately recalculated spreadsheet, not just Python formulas.
    small=load_workbook(OUT/'Risk_Formulas_Calculated.xlsx',read_only=True,data_only=True)
    engine=frame(small,'Risk_Excel_Formulas'); replay=frame(book,'Risk_Trade_Replay')
    assert engine.Date.tolist()==replay.Date.tolist()
    errors={}
    for calc,saved in [('Probability_formula','Loss_probability'),('Downside_formula','Expected_downside'),
                       ('Risk_formula','Risk_score'),('Threshold_formula','Risk_threshold')]:
        actual=pd.to_numeric(engine[calc],errors='raise').to_numpy(float); expected=replay[saved].to_numpy(float)
        np.testing.assert_allclose(actual,expected,atol=1e-9,rtol=0)
        errors[calc]=float(np.max(np.abs(actual-expected)))
    np.testing.assert_array_equal(engine.Dicey_formula.to_numpy(bool),replay.Dicey.to_numpy(bool))
    # Same inputs/coefficients as the master; the engine must not use stale data.
    for sheet in ('Risk_Inputs','Risk_Coefs','Risk_Rank_Keys'):
        original=frame(book,sheet); calculated=frame(small,sheet)
        assert original.columns.tolist()==calculated.columns.tolist()
        pd.testing.assert_frame_equal(original,calculated,check_dtype=False,atol=1e-12,rtol=1e-12)
    sheet_count=len(book.sheetnames); book.close(); small.close()
    result={'Status':'PASS','Iterations':len(registry),'Policies':count,'Period_rows_checked':periods,
        'Sheets':sheet_count,'Spreadsheet_engine':'LibreOffice headless; stored cached formula results verified',
        'Spreadsheet_risk_rows':len(engine),'Maximum_formula_errors':errors,
        'Checks':['input/output hashes','every daily equity/DD','trade totals and execution counts',
            'every quarterly/yearly/monthly/FY net and reset DD','actual spreadsheet formulas and flags',
            'spreadsheet inputs and coefficients agree with master']}
    (OUT/'verification.json').write_text(json.dumps(result,indent=2)+'\n')
    manifest['Input_sha256'][str(Path(__file__).relative_to(ROOT))]=hashlib.sha256(Path(__file__).read_bytes()).hexdigest()
    manifest['Stage']='verified'
    manifest['Output_sha256']={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in OUT.iterdir() if p.is_file() and p.name!='manifest.json'}
    (OUT/'manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
    print(json.dumps(result,indent=2),flush=True)


if __name__=='__main__': main()
