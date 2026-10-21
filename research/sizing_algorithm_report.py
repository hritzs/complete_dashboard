"""Describe and replay archived allocation algorithms; never optimize/refit."""
import json
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parent
CATALOG=[
 ('A00','Flat benchmark','None','size=1','Baseline allocation; no prediction','None','No learned confidence','All iterations'),
 ('A01','Shrunk bucket adjustment','Bucket mean PnL, overall mean mu, count n','a=0.6*(bucket_mean-mu)*n/(n+40)/mu; size=clip(1+sum(a), bounds)','Relative historical size contribution','Training split or full fit, depending on iteration','Small bucket shrinkage is not causal validation','sizing_equation_v2.py; sizing_full_history.py'),
 ('A02','S1 additive bucket score','Existing entry parameters; quintiles / discrete levels','S1=sum((bucket_mean-mu)*n/(n+20))','Relative profit score','Split-trained earlier; full-history later','Not probability; correlated effects can add twice','sizing_v5_capped.py:score_s1'),
 ('A03','S2 linear ridge score','Standardized parameters and table indicators','b=(Z.T*Z+10I)^-1*Z.T*(PnL-mu); score=Z*b','Centered profit score; v4 adds mu','Split-trained earlier; full-history later','Linear model, not neural net or calibrated edge','sizing_v4_params.py; sizing_v5_capped.py:score_s2'),
 ('A04','Score tercile allocation','S1/S2 and fitted 1/3,2/3 quantiles','tier=searchsorted(quantiles,score,right); size=levels[tier]','Fixed tier size such as .75/1/1.15','Same reference sample as score','Hard cap/09:16 and guards depend on iteration','sizing_v5_capped.py:tier_sizes'),
 ('A05','Volatility / stop exposure scaling','IV, RV10/20, straddle/spot, cushion, stop bps','clip(median(x)/x,.5,1.15); cushion uses x/median(x); stop uses 14/SLbps','Inverse-risk or stop scaling','Full-history medians in iteration09','Not an expected-profit model','sizing_full_history.py:build_candidates'),
 ('A06','VRP / calendar buckets','IV/RV, IV-RV, RV trend; DTE/event/weekday/gap','Bucket adjustment A01; high VRP alternative cuts .5 above 2/3 quantile','Regime-dependent size adjustment','Training split or full fit','Correlation does not establish stable live edge','sizing_vrp.py; sizing_v3.py; sizing_full_history.py'),
 ('A07','LUT-boundary / isolation restriction','Cell coordinates, 2%/5% distance, neighbor YES/NO, veto/floor','Cut .75/.5 on defined borderline or <=1 YES neighbor','Fragile-cell flag and fixed restriction','Published LUT plus fixed distances','Published LUT itself used broader history','sizing_v6_borderline.py:flags'),
 ('A08','Realized drawdown / streak guard','Earlier flat-book or sized-book outcomes','3 losses ->.5; flat DD>100 ->.5; max(.5,1-DD/200); v5 sized DD>=50 removes boost','Exposure restriction from already-realized stress','Sequential, strictly preceding outcome','Different guards use different equity books; no forecast','sizing_v5_capped.py:guard; sizing_full_history.py'),
 ('A09','Logistic loss probability','Transformed entry inputs; target 1[PnL<0]','sigmoid(clip(beta0+X*beta,-25,25)); L2 lambda20','Estimated probability of negative net PnL','Full-history in11/12; prior-year in13','Not independently calibrated; risk does not determine size alone','sizing_drawdown_risk.py:risk_fit; higher_tax_sizing.py:fit_risk'),
 ('A10','Ridge expected downside','Same risk inputs; target max(-PnL,0)','max(0,gamma0+X*gamma); ridge lambda50, unpenalized intercept','Expected negative part, including zero on winners','Same risk fitting sample','Not conditional loss size or VaR','sizing_drawdown_risk.py:risk_fit'),
 ('A11','Instability percentile average','Underlying range, LTP-straddle range, spread','I=(rank(Urange)+rank(StrRange)+rank(Spread))/3; missing/NW=.5','Relative pre-entry instability','Full-history or prior-year references','Legacy second-bucket extractor includes entry second','entry_window_features.py; sizing_drawdown_risk.py'),
 ('A12','Composite dicey gate','A09/A10 percentile ranks and A11','risk=.4*rank(p)+.4*rank(down)+.2*I; dicey=risk>=Q75 OR missing CW features','Mandatory-risk-cut flag','Same risk/reference fitting sample','40/40/20 fixed choices, not learned weights or probability','sizing_drawdown_risk.py:risk_fit'),
 ('A13','Linear continuous allocation','Transformed sizing inputs; allocation coefficients b','raw_size=1+X*b','Unconstrained allocation before offset and clipping','Full-history numerical optimization','Sizing b differs from risk beta/gamma; not a prediction probability','sizing_normalized_optimization.py; sizing_drawdown_risk.py'),
 ('A14','Multi-start coordinate pattern search','Saved PnL, X, hard bounds and objective','Fixed-seed starts; bounded coordinate steps; retain best feasible observed objective minus 1e-5*mean(b^2)','Allocation coefficient vector for DD/profit/ratio objective','Complete-history optimization','Local numerical search, not global optimum or live evidence','sizing_normalized_optimization.py:optimize; sizing_drawdown_risk.py:search'),
 ('A15','Bounded mean-one projection','raw_size, lower/upper bounds, historical count N','size=clip(raw_size+offset,lower,upper); solve sum(size)=N by43 bisections','Feasible historical sizes with mean exactly1','Entire historical trade list','Cannot enforce exact unknown future mean without future information','sizing_normalized_optimization.py:project'),
 ('A16','Earlier S1/S2 conviction gate','Standardized S1/S2, risk flag, loss probability, stage','CW & !dicey & S1>=CW_Q67 & S2>=CW_median & loss<=book_median; eligible excludes09:16','Strong estimate and priority allocation bounds','Full-history in iteration12','Ordinary optimizer boosts are reallocations, not strong-conviction labels','sizing_conviction_comparison.py:conviction_flags/allocation_bounds'),
 ('A17','Higher-tax monotone maximum closure','Prior-year cell post-cost profit P, weighted count N, implications','weight=round(100*(P-N*tax)); maximum closure via min-cut / iterative Dinic; anchor count>=1; stage cascade','Boolean higher-hurdle CW cell membership','Prior years; extra tax +2/+4','Tax is a selection hurdle, not deducted PnL or calibrated probability','lut_tax_closure.py:solve; higher_tax_sizing.py:fit_masks'),
 ('A18','Annual risk-first sizing branch','Warmup flag, dicey, CW higher-tax mask, fixed boost','warmup ->1; else dicey ->.85; else stronger CW ->1.05/1.10; else1','Final modest boost/cut; risk overrides conviction','Prior-year model and masks;2021 warmup','09:16 boost allowed only under original13 plus15/17; no future normalization; no candidate adopted','higher_tax_sizing.py:assign_sizes'),
 ('A19','Tick percentile / residual entry filters','Strict pre-entry premium rank, residual, spread/range/change','Eight fixed skip/delay rules; see TICK_ENTRY_SPEC.md','Entry rejection or one-minute delay, not size boost','Within-day prefix; prior-year across-day cutoffs','None met profit+ratio qualification; not executable stat arb','tick_entry_features.py; tick_entry_eval.py'),
 ('A20','Relaxed later entry','Time cutoff and YES-neighbor support of09:20+ NO cell','R1/R2/R3/R4 fixed windows and support; entry size1 or.75','Alternative entry set and timing','Fixed candidate evaluation on saved counterfactuals','Not simply sizing original entry; legacy LUT bias remains','entry_relax_eval.py'),
 ('A21','Annual frozen allocation rerun','Prior-year S1/S2, numeric scalers, risk models/references, allocation b and training offset','Fit prior years only; size=clip(1+X*b+frozen_offset,.75,cap); cap1.15, early CW1, WF11 dicey.85','Forward continuous size; actual future mean is reported','2021 warmup; annual expanding fit2022–2026; no old full-history weights','Training mean1 uses breakpoint projection; never normalize scored year; strict tick prefixes; legacy base LUT/fill bias remains','walkforward_sizing.py; WALKFORWARD_SIZING_SPEC.md'),
 ('A22','Monte Carlo scenario sizing','Annual prior-trained risk; stronger-tax masks at strict-prefix CW cells; fixed cut/boost grid','warmup1; benchmark=fixed size; otherwise dicey=cut; else tax-strong=boost; else1; no normalization','Fixed signal-based allocation and paired scenario profit/DD distributions','2022–2024 selection;2025+ evaluation;5000x150 trades/scenario','Six empirical/stress scenarios; not future forecasts, dynamic online refits or new edge; legacy LUT/fill/cache limitations remain','monte_carlo_sizing.py; MONTE_CARLO_SIZING_SPEC.md'),
 ('A23','Nested chronological sizing and placement','1944 risk/tax/boost/profit-score/IV/DD combinations; earlier inner train/validation','raw=risk?cut:(strong?boost:1); size=min(DDcap,strong?1.15:1,clip(raw*vol,.75,1.15)); flat fallback explicit','Parameters frozen before scored year; allocation-placement diagnostic at identical sizes','2021 model warmup;2022 calibration; outer2023+; prior-only60/40 inner split','Risk flags cannot be removed by admitted policies; fallback0 is baseline control; random permutations are diagnostics, not risk-gated policies; legacy base limitations','nested_sizing.py; NESTED_SIZING_SPEC.md'),
 ('A24','Matched execution-version replay','Identical frozen dates, daily inputs, events and latest LUTs; old and latest source snapshots','size=1; PnL comes from independent engine replay, never archived_PnL*size','Entry/exit/hedging implementation changes, separated from input-vintage changes','Fixed source-version comparison; no new model training or parameter selection','Earlier sizing studies use archived outcomes; option/futures fill and LUT assumptions remain; NW reported gross embeds costs','compare_latest_logic.py; LOGIC_COMPARISON_SPEC.md'),
 ('A25','Parallel directional long options','Past-only opening-range breakout and five-minute efficiency; optional lower AskIV percentile','Buy call/put,2:1 tilted straddle or balanced control; premium budget100 at entry; exit15min or30% premium stop; no futures hedge','Separate long-gamma/vega sleeve and actual GS-plus-sleeve daily portfolio','Fixed six-policy screen; no fitted model, tuning or historical mean normalization','Extra capital, fractional quantities, modeled atomic top-quote fills/fees and timestamp availability; no equal-risk or established-alpha claim','complement_strategy.py; COMPLEMENT_STRATEGY_SPEC.md')
]


def families(iteration,name):
    if name=='Flat size 1': return 'A00'
    if iteration=='20': return 'A25'
    if iteration=='19': return 'A24'
    if iteration=='18': return 'A09,A10,A11,A12,A17,A23'
    if iteration=='17': return 'A22' if name.startswith('Fixed') else 'A09,A10,A11,A12,A22'+(',A17' if name.startswith('Tax+') else '')
    if iteration=='15': return 'A02,A03,A13,A14,A21'
    if iteration=='16': return 'A02,A03,A09,A10,A11,A12,A13,A14,A21'
    if iteration=='13': return 'A09,A10,A11,A12,A17,A18'
    if iteration=='14': return 'A19'
    if iteration=='12': return 'A09,A10,A11,A12,A13,A14,A15'+('' if name.startswith('Previous') else ',A16')
    if iteration=='11': return 'A13,A14,A15' if name.startswith('Previous') else 'A09,A10,A11,A12,A13,A14,A15'
    if iteration=='10': return ('A13,A14,A15' if name.startswith('Optimized:') else families('09',name.removeprefix('Normalized: '))+',A15')
    if iteration=='09':
        if name.startswith('S1'): return 'A02,A04'+(',A08' if 'guard' in name else '')
        if name.startswith('S2'): return 'A03,A04'+(',A08' if 'guard' in name else '')
        if name.startswith(('Vol/','Stop:')): return 'A05'
        if name.startswith(('VRP','Calendar')): return 'A01,A06'
        if name.startswith(('Margin','Isolated')): return 'A07'
        if name in ('3 losses cut','DD100 cut','DD gradual cut'): return 'A08'
        return 'A01'
    return {'01':'A01','02':'A01','03':'A06','04':'A01,A06','05':'A02,A03,A04','06':'A02,A03,A04,A08','07':'A07','08':'A20'}[iteration]


def add_sheets(book,table,policy_map):
    paths={Path(__file__),ROOT/'SIZING_ALGORITHMS.md',ROOT/'research_iterations.json'}
    def load(relative,sheet=None):
        path=ROOT/'output'/relative; paths.add(path)
        return pd.read_excel(path,sheet_name=sheet) if path.suffix=='.xlsx' else pd.read_csv(path)
    table(book,'Algo_Catalog',pd.DataFrame(CATALOG,columns=['Algorithm_ID','Algorithm','Inputs','Formula','Determines','Fit_protocol','Limits_and_priority','Code_reference']))
    table(book,'Algo_Guide',pd.DataFrame({'Line':(ROOT/'SIZING_ALGORITHMS.md').read_text().splitlines()}))
    changes=[]
    for c in json.loads((ROOT/'research_iterations.json').read_text()):
        change={'09':'Full-history fit; global1.15 and09:16<=1; no mean rescaling',
            '10':'Historical mean-one projection; floor.75, cap1.15;09:16<=1',
            '11':'Mandatory dicey ceilings.90/.85/.75; three constrained objectives each',
            '12':'S1/S2 priority gate; dicey<=.85; ordinary<=1.10; early<=1; Gentle infeasible',
            '13':'Prior-year risk+tax gates;09:16 boosts allowed; risk.85 overrides; no future balancing',
            '14':'Strict tick prefixes; entry skip/delay; fixed multiplier1',
            '15':'Rebuild strict intraday inputs; all score/scaler/optimizer fits prior-year only; freeze training offset; actual mean',
            '16':'Same annual frozen protocol plus prior-year risk/percentile gates; dicey<=.85; explicit risk-only fallback',
            '17':'No normalization; paired Monte Carlo scenario paths; training-selected fixed sizing branches; exposure benchmarks separated',
            '18':'Nested earlier-data parameter selection; no normalization; IV/DD guards and stricter risk gates; identical-size placement nulls; explicit baseline fallbacks',
            '19':'Independent old/latest flat engine replays with identical current inputs and tables; archived vintage separate; actual entry/exit metadata, not sized old PnL',
            '20':'Separate premium-budgeted directional long-convexity sleeve alongside latest GS; six fixed policies and actual combined paths; flat benchmark changes to latest GS; no fitted alpha claim'}.get(c['id'],'Earlier split-based research; original restrictions and precise parameters in source tables')
        changes.append({'Iteration':c['id'],'Name':c['name'],'Change':change,'Protocol':c['protocol'],'Coverage':c['coverage']})
    table(book,'Algo_Changes',pd.DataFrame(changes))
    for row in policy_map: row['Algorithms']=families(row['Iteration'],row['Policy_name'])
    table(book,'Algo_Policy_Map',pd.DataFrame(policy_map))
    priorities=[
        ('Annual13',1,'Original base entry required','Base LUT and production gates stay unchanged; taxes13/5/2/2/6'),
        ('Annual13',2,'2021 warmup','size1; no annual learned estimate'),
        ('Annual13',3,'Dicey','size.85; dominates strong tax gate; NW may also be risk cut'),
        ('Annual13',4,'Non-dicey CW passes higher tax','boost1.05 or1.10;09:16 allowed after base13 plus15/17'),
        ('Annual13',5,'Otherwise','size1; NW never receives CW tax boost'),
        ('Archived12',1,'Dicey','upper.85; overrides strong; lower.75'),
        ('Archived12',2,'09:16 CW','upper1; strong estimate cannot override this archived restriction'),
        ('Archived12',3,'Eligible strong','Modest priority bounds1.10–1.15'),
        ('Archived12',4,'Ordinary','bounds.75–1.10; a boost may be normalization allocation'),
        ('Archived12',5,'Feasibility and projection','Require sum(lower)<=N<=sum(upper); project within bounds to mean1'),
        ('Archived10/11/12',6,'Qualification','Historical objective/constraints; not an approved live policy')]
    table(book,'Algo_Priority',pd.DataFrame(priorities,columns=['Version','Order','Condition','Action']))
    conviction=load('conviction_sizing/trade_comparison.csv'); source=load('conviction_sizing/conviction_trade_comparison.xlsx')
    inputs=source['Model_inputs'].set_index('Date'); gates=source['Gate_thresholds'].iloc[0]
    cw=conviction.Leg.eq('CW').to_numpy(); early=cw&conviction.Stage.eq('09:16').to_numpy(); dicey=conviction.Dicey.to_numpy(bool)
    s1=inputs.loc[conviction.Date,'S1_fullfit'].to_numpy(); s2=inputs.loc[conviction.Date,'S2_fullfit'].to_numpy()
    a=s1>=gates.S1_top_tercile; b=s2>=gates.S2_median; low=conviction.Loss_probability.to_numpy()<=gates.Loss_probability_median
    strong=cw&~dicey&a&b&low; eligible=strong&~early
    np.testing.assert_array_equal(strong,conviction.Strong_estimated_edge); np.testing.assert_array_equal(eligible,conviction.Boost_eligible)
    gate=pd.DataFrame({'Date':conviction.Date,'Leg':conviction.Leg,'Stage':conviction.Stage,'Dicey':dicey,'S1':s1,'S1_cutoff':gates.S1_top_tercile,
        'S1_pass':a,'S2':s2,'S2_cutoff':gates.S2_median,'S2_pass':b,'Loss_probability':conviction.Loss_probability,
        'Loss_cutoff':gates.Loss_probability_median,'Low_loss_pass':low,'Strong':strong,'Early0916':early,'Boost_eligible':eligible,
        'New_size':conviction.New_size,'Saved_reason':conviction.Reason})
    table(book,'Algo_Conviction',gate); table(book,'Algo_Profile_Feasibility',source['Profile_feasibility'])
    annual=load('higher_tax_sizing/trade_comparison.csv'); annual_rows=[]
    for delta in (2,4):
        for boost in (1.05,1.10):
            name=f'Tax+{delta} boost{boost:g}'; warm=annual.Warmup.to_numpy(bool)
            risk=annual.Dicey.to_numpy(bool); high=annual[f'Strong_plus{delta}'].fillna(False).to_numpy(bool)
            iscw=annual.Leg.eq('CW').to_numpy(); predicted=np.where(warm,1.,np.where(risk,.85,np.where(iscw&high,boost,1.)))
            np.testing.assert_allclose(predicted,annual[name+'_size'],atol=1e-12)
            for i,r in annual.iterrows():
                rr=len(annual_rows)+2
                reason='Warmup size1' if warm[i] else 'Dicey overrides conviction' if risk[i] else 'CW higher-tax conviction' if iscw[i]&high[i] else 'Ordinary size1; no CW tax boost for NW'
                annual_rows.append({'Date':r.Date,'Candidate':name,'Leg':r.Leg,'Stage':r.Stage,'Warmup':bool(warm[i]),'Dicey':bool(risk[i]),
                    'Higher_tax_pass':bool(high[i]),'Boost':boost,'Base_tax':r.Base_tax,'Higher_tax':r.Base_tax+delta if pd.notna(r.Base_tax) else np.nan,
                    'Replayed_size':predicted[i],'Saved_size':r[name+'_size'],'Reason':reason,
                    'Size_Excel':f'=IF(E{rr},1,IF(F{rr},0.85,IF(AND(C{rr}="CW",G{rr}),H{rr},1)))',
                    'Excel_error':f'=N{rr}-L{rr}'})
    table(book,'Algo_Annual',pd.DataFrame(annual_rows))
    # Replay all saved numerical optimized policies, without searching again.
    dd=load('drawdown_risk_sizing/drawdown_risk_sizing.xlsx')
    all_inputs=dd['Model_inputs']; names=[n for n in all_inputs if n!='Date']; matrix=all_inputs[names].to_numpy(float)
    traded=dd['Daily_risk']; td=traded[traded.Leg.isin(['CW','NW'])].reset_index(drop=True)
    assert td.Date.tolist()==all_inputs.Date.tolist()==conviction.Date.tolist()
    table(book,'Algo_Opt_Inputs',all_inputs)
    replay=[]; coefficient_rows=[]; checks=[]
    for iteration,relative in [('10','normalized_sizing/normalized_sizing.xlsx'),('11','drawdown_risk_sizing/drawdown_risk_sizing.xlsx'),('12','conviction_sizing/conviction_trade_comparison.xlsx')]:
        saved=load(relative); coef=saved['Coefficients']; key='Objective' if iteration=='10' else 'Candidate'
        daily=saved['Daily_sizes'].set_index('Date'); policy_names=list(saved['Daily_PnL'].columns.drop('Date'))
        for candidate,group in coef.groupby(key,sort=False):
            weights=group.set_index('Feature').Coefficient
            if not set(weights.index)<=set(names): raise ValueError('Unknown optimization feature')
            beta=np.array([weights.get(n,0.) for n in names]); offset_col='Normalization_offset' if iteration=='10' else 'Offset'
            assert group[offset_col].nunique()==1
            offset=float(group[offset_col].iloc[0]); raw=1+matrix@beta
            lower=np.full(len(td),.75); upper=np.where(td.Leg.eq('CW')&td.Stage.eq('09:16'),1.,1.15)
            if iteration=='11':
                ceiling=float(candidate.split('<=')[1].split(':')[0]); upper[dicey]=np.minimum(upper[dicey],ceiling)
            elif iteration=='12':
                lower=np.where(eligible,1.10,.75); upper=np.where(eligible,1.15,1.10)
                upper[early]=np.minimum(upper[early],1.); upper[dicey]=np.minimum(upper[dicey],.85)
            predicted=np.clip(raw+offset,lower,upper); name='Optimized: '+candidate if iteration=='10' else candidate
            archived=daily.loc[td.Date,name].to_numpy(float)
            np.testing.assert_allclose(predicted,archived,atol=1e-10,rtol=0)
            assert abs(predicted.mean()-1)<1e-10
            policy_id=f'{iteration}_P{policy_names.index(name):02d}'; cr=len(coefficient_rows)+2
            coefficient_rows.append({'Policy':policy_id,**{n:beta[j] for j,n in enumerate(names)},'Offset':offset})
            last=get_column_letter(len(names)+1)
            checks.append({'Policy':policy_id,'Name':name,'Mean_size':float(predicted.mean()),'Maximum_error':float(np.max(abs(predicted-archived))),
                'Lower_sum':float(lower.sum()),'Upper_sum':float(upper.sum()),'Target':len(td)})
            for i,r in td.iterrows():
                rr=len(replay)+2; ir=i+2; shifted=raw[i]+offset
                binding='floor' if shifted<lower[i] else 'ceiling' if shifted>upper[i] else 'interior'
                reason='Mandatory dicey upper bound' if iteration in ('11','12') and dicey[i] else 'Strong-edge priority bounds' if iteration=='12' and eligible[i] else 'Archived09:16 cap' if early[i] else 'Optimized allocation/budget balancing; not a strong-edge label'
                replay.append({'Date':r.Date,'Policy':policy_id,'Candidate':name,'Dicey':bool(dicey[i]),'Strong_eligible':bool(eligible[i]),
                    'Early0916':bool(early[i]),'Raw_size':raw[i],'Offset':offset,'Lower':lower[i],'Upper':upper[i],
                    'Replayed_size':predicted[i],'Saved_size':archived[i],'Replay_error':predicted[i]-archived[i],
                    'Binding':binding,'Reason':reason,
                    'Raw_Excel':f'=1+SUMPRODUCT(Algo_Opt_Inputs!B{ir}:{last}{ir},Algo_Opt_Coefs!$B${cr}:${last}${cr})',
                    'Size_Excel':f'=MAX(I{rr},MIN(J{rr},P{rr}+H{rr}))','Excel_error':f'=Q{rr}-L{rr}'})
    table(book,'Algo_Opt_Coefs',pd.DataFrame(coefficient_rows)); table(book,'Algo_Opt_Replay',pd.DataFrame(replay))
    table(book,'Algo_Opt_Checks',pd.DataFrame(checks))
    # Conditional walk-forward reruns: saved fold coefficients, never a refit.
    wf=load('walkforward_sizing/walkforward_comparison.xlsx')
    wf_inputs=wf['Model_inputs']; wf_names=list(dict.fromkeys(wf['Coefficients'].Feature))
    wf_coefs=[];wf_rows=[];wf_input_rows=[]
    wf_trade=wf['Trade_comparison'].set_index('Date')
    for (year,candidate),group in wf['Coefficients'].groupby(['Year','Candidate'],sort=False):
        policy_id='15_P01' if candidate.startswith('WF10') else '16_P01'
        beta=group.set_index('Feature').Coefficient.reindex(wf_names,fill_value=0).to_numpy(float)
        assert group.Frozen_offset.nunique()==1
        offset=float(group.Frozen_offset.iloc[0]);cr=len(wf_coefs)+2
        wf_coefs.append({'Year':int(year),'Policy':policy_id,**dict(zip(wf_names,beta)),'Offset':offset})
        selected=wf_inputs[wf_inputs.Year.eq(year)&wf_inputs.Policy.eq(candidate)]
        matrix=selected[wf_names].fillna(0).to_numpy(float);raw=1+matrix@beta
        last=get_column_letter(len(wf_names)+2)
        for i,date in enumerate(selected.Date):
            source=wf_trade.loc[date];early=source.Leg=='CW' and source.Stage=='09:16';dice=bool(source.Dicey)
            cap=min(1. if early else 1.15,.85 if policy_id=='16_P01' and dice else 1.15)
            predicted=float(np.clip(raw[i]+offset,.75,cap));archived=float(source[candidate+'_size'])
            np.testing.assert_allclose(predicted,archived,atol=1e-10,rtol=0)
            ir=len(wf_input_rows)+2;rr=len(wf_rows)+2
            wf_input_rows.append({'Date':date,'Policy':policy_id,**dict(zip(wf_names,matrix[i]))})
            wf_rows.append({'Date':date,'Year':int(year),'Policy':policy_id,'Candidate':candidate,'Dicey':dice,
                'Early0916':early,'Raw_size':raw[i],'Offset':offset,'Lower':.75,'Upper':cap,
                'Replayed_size':predicted,'Saved_size':archived,
                'Raw_Excel':f'=1+SUMPRODUCT(Algo_WF_Inputs!C{ir}:{last}{ir},Algo_WF_Coefs!$C${cr}:${last}${cr})',
                'Size_Excel':f'=MAX(I{rr},MIN(J{rr},M{rr}+H{rr}))','Excel_error':f'=N{rr}-L{rr}'})
    table(book,'Algo_WF_Inputs',pd.DataFrame(wf_input_rows));table(book,'Algo_WF_Coefs',pd.DataFrame(wf_coefs))
    table(book,'Algo_WF_Replay',pd.DataFrame(wf_rows));table(book,'Algo_WF_Folds',wf['Training_cutoffs'])
    paths.update([ROOT/'WALKFORWARD_SIZING_SPEC.md',ROOT/'walkforward_sizing.py',ROOT/'audit_walkforward_sizing.py'])
    mc=load('monte_carlo_sizing/monte_carlo_sizing.xlsx');mc_rows=[]
    mc_trades=mc['Trade_comparison']
    for rule in mc['Policy_parameters'].itertuples():
        for row in mc_trades.to_dict('records'):
            rr=len(mc_rows)+2;warm=bool(row['Warmup']);dice=bool(row['Dicey'])
            mc_strong=bool(row[f'Strong_tax_plus{int(rule.Extra_tax)}']) if pd.notna(rule.Extra_tax) else False
            mc_rows.append({'Date':row['Date'],'Candidate':rule.Candidate,'Warmup':warm,'Dicey':dice,'Strong':mc_strong,
                'Risk_cut':float(rule.Risk_cut) if pd.notna(rule.Risk_cut) else 1.,'Boost':float(rule.Boost),'Family':rule.Family,
                'Saved_size':row[rule.Candidate+'_size'],
                'Size_Excel':f'=IF(C{rr},1,IF(H{rr}="Exposure benchmark",G{rr},IF(D{rr},F{rr},IF(E{rr},G{rr},1))))',
                'Excel_error':f'=J{rr}-I{rr}'})
    table(book,'Algo_MC_Replay',pd.DataFrame(mc_rows));table(book,'Algo_MC_Parameters',mc['Policy_parameters'])
    table(book,'Algo_MC_Selection',mc['Frozen_selection'])
    paths.update([ROOT/'MONTE_CARLO_SIZING_SPEC.md',ROOT/'monte_carlo_sizing.py'])
    nested=load('nested_sizing/nested_sizing.xlsx');nr=nested['Formula_replay'].copy()
    formulas=[]
    for i in range(len(nr)):
        r=i+2
        formula=f'=IF(E{r},1,MIN(K{r},IF(G{r},1.15,1),MAX(0.75,MIN(1.15,IF(F{r},H{r},IF(G{r},I{r},1))*J{r}))))'
        assert formula.count('(')==formula.count(')');formulas.append(formula)
    nr['Size_Excel']=formulas;nr['Excel_error']=[f'=M{i+2}-L{i+2}' for i in range(len(nr))]
    table(book,'Algo_Nested_Replay',nr);table(book,'Algo_Nested_Selection',nested['Frozen_parameters'])
    paths.update([ROOT/'NESTED_SIZING_SPEC.md',ROOT/'nested_sizing.py'])
    return paths,{'Algorithms':len(CATALOG),'Mapped_policies':len(policy_map),'Conviction_rows':len(gate),
        'Strong':int(strong.sum()),'Eligible':int(eligible.sum()),'Annual_sizing_rows':len(annual_rows),
        'Optimized_policies':len(checks),'Optimized_trade_rows':len(replay),'Maximum_size_replay_error':max(c['Maximum_error'] for c in checks),
        'Walkforward_sizing_rows':len(wf_rows),'Walkforward_folds':len(wf_coefs),'Monte_Carlo_sizing_rows':len(mc_rows),
        'Nested_sizing_formula_rows':len(nr)}
