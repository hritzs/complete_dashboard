"""
Summary workbook of the GammaShort research results saved on the share.

Reads (never writes) three files in \\\\172.16.1.85\\Shared\\GammaShortData\\GammaShort:
  combined_book_daily.csv  daily PnL of the combined CW + NW book
  ps1_daily_walk.csv       the current-week leg per day (entry minute, DTE)
  nw_sim_daily.csv         the next-week leg on the days it was simulated
and writes research/output/GammaShort_history_summary.xlsx.

PnL in the files is "normalised Rs": index points x 25000 / entry underlying,
i.e. the result of one trade sized at Rs 25,000 of notional. Rupees for a real
size are therefore  normalised x (notional / 25000).  The workbook keeps the
notional in one cell (Summary!B4) and every rupee figure is a formula on it.

Standard library + openpyxl only.   python3 research/make_history_summary.py
"""
import csv
import os
from collections import OrderedDict
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SRC = os.environ.get("GS_RESEARCH_DIR", "/mnt/shared/GammaShortData/GammaShort")
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "output", "GammaShort_history_summary.xlsx")
UNIT_NOTIONAL = 25000.0
DEFAULT_LOTS, LOT_SIZE, REF_INDEX = 77, 65, 22500.0     # the platform's live size, for the default


def d(s):  # DDMMYYYY
    return date(int(s[4:]), int(s[2:4]), int(s[:2]))


def fy(dt):
    y = dt.year if dt.month >= 4 else dt.year - 1
    return f"FY{str(y)[-2:]}-{str(y + 1)[-2:]}"


def read(name):
    with open(os.path.join(SRC, name), newline="") as f:
        return list(csv.DictReader(f))


def num(x):
    try:
        v = float(x)
        return v if v == v else None
    except (TypeError, ValueError):
        return None


def stats(pnls):
    """Total, trades, wins, max drawdown (equity starts at 0), peak / trough index."""
    cum = hw = 0.0
    mdd = 0.0
    peak_i = trough_i = cur_peak = -1
    for i, p in enumerate(pnls):
        cum += p
        if cum > hw:
            hw, cur_peak = cum, i
        if hw - cum > mdd:
            mdd, peak_i, trough_i = hw - cum, cur_peak, i
    tr = [p for p in pnls if p != 0]
    wins = [p for p in tr if p > 0]
    losses = [p for p in tr if p < 0]
    return dict(total=sum(pnls), days=len(pnls), trades=len(tr), wins=len(wins),
                win_pct=100.0 * len(wins) / len(tr) if tr else 0.0,
                avg=sum(tr) / len(tr) if tr else 0.0,
                avg_win=sum(wins) / len(wins) if wins else 0.0,
                avg_loss=sum(losses) / len(losses) if losses else 0.0,
                best=max(pnls) if pnls else 0.0, worst=min(pnls) if pnls else 0.0,
                mdd=mdd, peak_i=peak_i, trough_i=trough_i)


def episodes(dates, pnls, top=15):
    """Drawdown episodes: peak -> trough -> recovery, deepest first."""
    out = []
    cum = hw = 0.0
    peak = None
    cur = None
    for dt, p in zip(dates, pnls):
        cum += p
        if cum >= hw:
            if cur:
                cur["recovered"] = dt
                out.append(cur)
                cur = None
            hw, peak = cum, dt
        else:
            if cur is None:
                cur = dict(peak=peak, trough=dt, depth=hw - cum, recovered=None)
            elif hw - cum > cur["depth"]:
                cur["trough"], cur["depth"] = dt, hw - cum
    if cur:
        out.append(cur)
    return sorted(out, key=lambda e: -e["depth"])[:top]


# ---------------------------------------------------------------- load
comb = [(d(r["date"]), float(r["pnl"])) for r in read("combined_book_daily.csv")]
comb.sort()
dates = [x[0] for x in comb]
pnls = [x[1] for x in comb]
walk = {d(r["date"]): r for r in read("ps1_daily_walk.csv")}
nwd = {d(r["date"]): r for r in read("nw_sim_daily.csv")}

S = stats(pnls)

# ---------------------------------------------------------------- workbook
wb = Workbook()
H = Font(bold=True)
FILL = PatternFill("solid", fgColor="DDEBF7")
INPUT = PatternFill("solid", fgColor="FFF2CC")
RS = '#,##0;[Red]-#,##0'
N1 = '#,##0.0;[Red]-#,##0.0'


def header(ws, row, cols):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill = H, FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---- Summary
ws = wb.active
ws.title = "Summary"
ws["A1"] = "GammaShort - history summary (research results on the share)"
ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = f"{dates[0]:%d-%b-%Y} to {dates[-1]:%d-%b-%Y}, {len(dates)} days. Combined book = current-week + next-week legs."
ws["A4"], ws["B4"] = "Notional per trade (Rs)  <- change this", DEFAULT_LOTS * LOT_SIZE * REF_INDEX
ws["B4"].fill, ws["B4"].number_format, ws["A4"].font = INPUT, RS, H
ws["C4"] = f"default = {DEFAULT_LOTS} lots x {LOT_SIZE} x index {REF_INDEX:,.0f}"
ws["A5"], ws["B5"] = "Rupees per 1 normalised Rs", "=B4/25000"
ws["B5"].number_format = '#,##0.00'
ws["C5"] = "the files are sized at Rs 25,000 notional per trade"
F = "Summary!$B$5"

header(ws, 7, ["Metric", "Normalised Rs", "Rupees at the notional above", "Note"])
rows = [
    ("Net PnL (profit)", S["total"], True, ""),
    ("Max drawdown", S["mdd"], True, f"{dates[S['peak_i']]:%d-%b-%Y} peak to {dates[S['trough_i']]:%d-%b-%Y} trough"),
    ("PnL / max drawdown", S["total"] / S["mdd"] if S["mdd"] else 0, False, "size-independent"),
    ("Days in the history", S["days"], False, ""),
    ("Days with a trade", S["trades"], False, "days with non-zero PnL"),
    ("Winning trades %", S["win_pct"], False, ""),
    ("Average per trade", S["avg"], True, ""),
    ("Average winning trade", S["avg_win"], True, ""),
    ("Average losing trade", S["avg_loss"], True, ""),
    ("Best day", S["best"], True, f"{dates[pnls.index(S['best'])]:%d-%b-%Y}"),
    ("Worst day", S["worst"], True, f"{dates[pnls.index(S['worst'])]:%d-%b-%Y}"),
    ("Average PnL per year", S["total"] / ((dates[-1] - dates[0]).days / 365.25), True, "calendar years in the history"),
]
for i, (k, v, rupee, note) in enumerate(rows, 8):
    ws.cell(row=i, column=1, value=k)
    c = ws.cell(row=i, column=2, value=round(v, 2))
    c.number_format = N1
    if rupee:
        r = ws.cell(row=i, column=3, value=f"=B{i}*$B$5")
        r.number_format = RS
    ws.cell(row=i, column=4, value=note)

r0 = 8 + len(rows) + 1
ws.cell(row=r0, column=1, value="Sizing table").font = H
header(ws, r0 + 1, ["Size", "Notional (Rs)", "Net PnL (Rs)", "Max drawdown (Rs)", "Worst day (Rs)", "Avg per trade (Rs)"])
sizes = [("1 lot (65)", 1), ("10 lots", 10), ("25 lots", 25), ("50 lots", 50), ("77 lots (live size)", 77),
         ("100 lots", 100), ("150 lots", 150)]
for i, (name, lots) in enumerate(sizes, r0 + 2):
    ws.cell(row=i, column=1, value=name)
    ws.cell(row=i, column=2, value=lots * LOT_SIZE * REF_INDEX).number_format = RS
    for col, val in ((3, S["total"]), (4, S["mdd"]), (5, S["worst"]), (6, S["avg"])):
        ws.cell(row=i, column=col, value=f"=B{i}/25000*{val:.6f}").number_format = RS
n0 = r0 + 2 + len(sizes) + 1
notes = [
    "How sizing works: every backtest trade is Rs 25,000 of notional, whatever the index level. Rupees = normalised x notional / 25,000.",
    f"Lots are converted at index {REF_INDEX:,.0f} (1 lot = Rs {LOT_SIZE * REF_INDEX:,.0f} notional). A fixed number of lots was a smaller notional when the index was lower,",
    "so 'lots' columns describe a constant-notional book equal to that many lots today, not a fixed lot count through history.",
    "Max drawdown is on the daily closing equity (starting at 0), the same definition as the backtest script.",
    "Costs are already inside these numbers (entry slippage and hedge cost as modelled by the research). Brokerage, STT and exchange charges are not.",
    "Source: combined_book_daily.csv dated 02-Sep-2026 on the share. It predates the 17-Sep 'no-roll' re-base and the round-52 script in this folder,",
    "whose header quotes 10,640.5 / MDD 144.4 over 1,379 days. This workbook reports the saved file, not a new run of the script.",
    "Daily sheet, Leg column: CW = the day's PnL equals the current-week file; NW = it differs and the day is in the next-week file;",
    "n/a = it differs and the saved files do not say which leg or entry produced it.",
]
for i, t in enumerate(notes, n0):
    ws.cell(row=i, column=1, value=t)
widths(ws, [42, 20, 28, 22, 18, 20])

# ---- By FY
ws = wb.create_sheet("By_FY")
header(ws, 1, ["Financial year", "Days", "Trades", "Win %", "Net PnL (norm)", "Max DD (norm)", "Worst day (norm)",
               "Net PnL (Rs)", "Max DD (Rs)", "Worst day (Rs)"])
byfy = OrderedDict()
for dt, p in comb:
    byfy.setdefault(fy(dt), []).append(p)
for i, (k, v) in enumerate(byfy.items(), 2):
    s = stats(v)
    vals = [k, s["days"], s["trades"], round(s["win_pct"], 1), round(s["total"], 1), round(s["mdd"], 1), round(s["worst"], 1)]
    for c, x in enumerate(vals, 1):
        ws.cell(row=i, column=c, value=x)
    for c, src in ((8, "E"), (9, "F"), (10, "G")):
        ws.cell(row=i, column=c, value=f"={src}{i}*{F}").number_format = RS
widths(ws, [16, 8, 8, 8, 16, 16, 16, 18, 18, 18])

# ---- By month
ws = wb.create_sheet("By_Month")
header(ws, 1, ["Month", "Trades", "Win %", "Net PnL (norm)", "Worst day (norm)", "Net PnL (Rs)", "Worst day (Rs)"])
bym = OrderedDict()
for dt, p in comb:
    bym.setdefault(dt.strftime("%Y-%m"), []).append(p)
for i, (k, v) in enumerate(bym.items(), 2):
    s = stats(v)
    for c, x in enumerate([k, s["trades"], round(s["win_pct"], 1), round(s["total"], 1), round(s["worst"], 1)], 1):
        ws.cell(row=i, column=c, value=x)
    ws.cell(row=i, column=6, value=f"=D{i}*{F}").number_format = RS
    ws.cell(row=i, column=7, value=f"=E{i}*{F}").number_format = RS
widths(ws, [12, 8, 8, 16, 16, 18, 18])

# ---- Current-week leg by DTE and entry stage (ps1_daily_walk)
ws = wb.create_sheet("CW_by_DTE_and_entry")
header(ws, 1, ["Group", "Trades", "Win %", "Net PnL (norm)", "Avg per trade (norm)", "Worst (norm)", "Net PnL (Rs)", "Worst (Rs)"])
fired = [(dt, r) for dt, r in sorted(walk.items()) if r.get("fired") == "True" and num(r.get("cw")) is not None]


def stage(e):
    e = int(float(e))
    return {916: "09:16", 917: "09:17", 918: "09:18", 919: "09:19"}.get(e, "09:20-09:59" if e < 1000 else "10:00 onwards")


groups = OrderedDict()
for dt, r in fired:
    groups.setdefault(f"DTE {int(float(r['dte'])) + 1}", []).append(float(r["cw"]))
groups = OrderedDict(sorted(groups.items()))
st = OrderedDict((k, []) for k in ["09:16", "09:17", "09:18", "09:19", "09:20-09:59", "10:00 onwards"])
for dt, r in fired:
    st[stage(r["entry"])].append(float(r["cw"]))
row = 2
for title, grp in (("By days to expiry (1 = expiry day)", groups), ("By entry time", st)):
    ws.cell(row=row, column=1, value=title).font = H
    row += 1
    for k, v in grp.items():
        s = stats(v)
        for c, x in enumerate([k, s["trades"], round(s["win_pct"], 1), round(s["total"], 1), round(s["avg"], 2), round(s["worst"], 1)], 1):
            ws.cell(row=row, column=c, value=x)
        ws.cell(row=row, column=7, value=f"=D{row}*{F}").number_format = RS
        ws.cell(row=row, column=8, value=f"=F{row}*{F}").number_format = RS
        row += 1
    row += 1
ws.cell(row=row, column=1, value="Source: ps1_daily_walk.csv (current-week leg only, before the next-week leg replaces some days).")
widths(ws, [34, 8, 8, 16, 20, 14, 18, 16])

# ---- Drawdowns
ws = wb.create_sheet("Drawdowns")
header(ws, 1, ["Rank", "Peak date", "Trough date", "Recovered", "Days peak to trough", "Depth (norm)", "Depth (Rs)"])
for i, e in enumerate(episodes(dates, pnls), 2):
    vals = [i - 1, e["peak"], e["trough"], e["recovered"] or "not yet", (e["trough"] - e["peak"]).days if e["peak"] else "", round(e["depth"], 1)]
    for c, x in enumerate(vals, 1):
        ws.cell(row=i, column=c, value=x)
    ws.cell(row=i, column=7, value=f"=F{i}*{F}").number_format = RS
widths(ws, [6, 14, 14, 14, 20, 14, 18])

# ---- Daily (complete history)
ws = wb.create_sheet("Daily")
header(ws, 1, ["Date", "FY", "Leg", "CW entry", "DTE", "PnL (norm)", "Cum PnL (norm)", "Drawdown (norm)",
               "PnL (Rs)", "Cum PnL (Rs)", "Drawdown (Rs)"])
cum = hw = 0.0
for i, (dt, p) in enumerate(comb, 2):
    cum += p
    hw = max(hw, cum)
    w = walk.get(dt, {})
    n = nwd.get(dt, {})
    cw = num(w.get("cw"))
    leg = ""
    if p != 0:
        # CW  = equals the current-week walk file's result for the day
        # NW  = differs from it and the day is in the next-week file
        # n/a = differs from it and is not in the next-week file (the saved
        #       files do not say which leg or entry produced it)
        leg = "CW" if (cw is not None and abs(cw - p) < 1e-6) else ("NW" if n else "n/a")
    entry = w.get("entry", "")
    entry = f"{int(float(entry)) // 100:02d}:{int(float(entry)) % 100:02d}" if leg == "CW" and entry not in ("", "nan") else ""
    dte = int(float(w["dte"])) + 1 if w.get("dte") not in (None, "", "nan") else ""
    for c, x in enumerate([dt, fy(dt), leg, entry, dte, round(p, 2), round(cum, 2), round(cum - hw, 2)], 1):
        ws.cell(row=i, column=c, value=x)
    for c, src in ((9, "F"), (10, "G"), (11, "H")):
        ws.cell(row=i, column=c, value=f"={src}{i}*{F}").number_format = RS
ws.freeze_panes = "A2"
widths(ws, [12, 10, 6, 10, 6, 12, 16, 16, 16, 18, 16])

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)

legs = {}
for r in ws.iter_rows(min_row=2, min_col=3, max_col=3, values_only=True):
    legs[r[0]] = legs.get(r[0], 0) + 1
print(f"saved {OUT}")
print(f"days {S['days']}  trades {S['trades']}  net {S['total']:.1f}  mdd {S['mdd']:.1f} "
      f"({dates[S['peak_i']]} -> {dates[S['trough_i']]})  win {S['win_pct']:.1f}%  avg {S['avg']:.2f}  worst {S['worst']:.1f}")
print("legs:", legs)
for k, v in byfy.items():
    s = stats(v)
    print(f"  {k}: pnl {s['total']:8.1f}  mdd {s['mdd']:6.1f}  trades {s['trades']}")
