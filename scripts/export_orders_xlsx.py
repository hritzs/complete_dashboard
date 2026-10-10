#!/usr/bin/env python3
"""Export a day's order log and trade book to Excel.

Source of truth: the broker's own order / trade pushes (Iris), as recorded in
6_reconciler.log, plus the platform's order and exit events from
4_execution.log.

  scripts/export_orders_xlsx.py [LOG_DIR] [OUT.xlsx]

LOG_DIR defaults to ./logs (today); a past day is logs/archive/YYYY-MM-DD.
OUT defaults to exports/orders_<day>_<HHMMSS>.xlsx.

Sheets: Summary, Trade Book, Orders, Order Log, Platform Events.
"""
import json
import os
import re
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
log_dir = sys.argv[1] if len(sys.argv) > 1 else os.path.join(BASE, "logs")
SIDE = {"1": "BUY", "2": "SELL"}
OTYPE = {"1": "LIMIT", "2": "MARKET"}
LINE = re.compile(r"^(\d{4})/(\d{2})/(\d{2}) (\d{2}:\d{2}:\d{2}) (.*)$")
RX = re.compile(r"\[GREEKSOFT IRIS RX\] streaming_type=(\w+) service=\w+ raw=(\{.*\})\s*$")
EVENT = re.compile(
    r"\[MANUAL ORDER\]|\[GREEKSOFT ORDER\] sending|BUILD (submitted|outcome|verification)|\[BUILD-CHASE\]|\[PARTIAL\]|"
    r"SQF reconciliation|Square-off|\[SQUARE-OFF-ALL\]|\[MTM-EXIT\]|\[STRADDLE-EXIT\]|\[RISK\]|\[PORTFOLIO-MTM\]|"
    r"\[WINGS\]|\[HEDGE|\[MODIFY\]|\[LUT-BUILD\]|\[SBUILD-LIVE\]|\[MOCK\]|realized PnL stored|DeployStraddle"
)


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def lines(name):
    path = os.path.join(log_dir, name)
    if not os.path.exists(path):
        return
    with open(path, errors="replace") as f:
        for ln in f:
            m = LINE.match(ln.rstrip("\n"))
            if m:
                yield f"{m.group(1)}-{m.group(2)}-{m.group(3)}", m.group(4), m.group(5)


day = ""
order_log, trades, orders = [], [], {}
for d, t, rest in lines("6_reconciler.log"):
    m = RX.search(rest)
    if not m or m.group(1) in ("HeartBeat", "LicenseResponse"):
        continue
    try:
        x = json.loads(m.group(2))["response"]["data"]
    except (ValueError, KeyError):
        continue
    day = day or d
    oid = x.get("gorderid", "")
    side = SIDE.get(str(x.get("side")), str(x.get("side", "")))
    typ = "MARKET" if str(x.get("is_MarketOrder")) == "1" else OTYPE.get(str(x.get("order_type")), "")
    reason = (x.get("reason") or x.get("remarks") or "").strip()
    row = [d, t, m.group(1), oid, x.get("eorderid", ""), x.get("symbol", ""), side, num(x.get("qty")),
           num(x.get("price")), typ, x.get("order_status", ""), num(x.get("pending_qty")),
           num(x.get("qty_filled_today")), reason, x.get("userTag", "")]
    order_log.append(row)
    o = orders.setdefault(oid, {"date": d, "first": t, "symbol": "", "side": "", "qty": None, "limit": None,
                                "type": "", "status": "", "filled": 0.0, "value": 0.0, "reason": "", "tag": ""})
    o["last"] = t
    for k, v in (("symbol", x.get("symbol", "")), ("side", side), ("type", typ), ("tag", x.get("userTag", ""))):
        if v:
            o[k] = v
    if num(x.get("qty")) is not None:
        o["qty"] = num(x.get("qty"))
    if num(x.get("price")) is not None and m.group(1) != "TradeResponse":
        o["limit"] = num(x.get("price"))
    if x.get("order_status"):
        o["status"] = x["order_status"]
    if reason:
        o["reason"] = reason
    if m.group(1) == "TradeResponse":
        q, p = num(x.get("traded_qty")) or 0, num(x.get("traded_price")) or 0
        o["filled"] += q
        o["value"] += q * p
        trades.append([d, t, x.get("tradeid", ""), oid, x.get("eorderid", ""), x.get("symbol", ""), side, q, p,
                       None, num(x.get("price")), x.get("userTag", "")])

events = []
for d, t, rest in lines("4_execution.log"):
    if EVENT.search(rest):
        day = day or d
        events.append([d, t, rest[:1500]])

day = day or datetime.now().strftime("%Y-%m-%d")
out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(
    BASE, "exports", f"orders_{day}_{datetime.now().strftime('%H%M%S')}.xlsx")
os.makedirs(os.path.dirname(out), exist_ok=True)

wb = Workbook()
FONT, BOLD = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEAD = PatternFill("solid", fgColor="1F3864")


def sheet(ws, headers, rows, widths, money=()):
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font, c.fill, c.alignment = BOLD, HEAD, Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = FONT
            if c.column in money:
                c.number_format = "#,##0.00"
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions


# Trade Book (Value = Qty x Price as a formula).
tb = wb.active
tb.title = "Trade Book"
for i, r in enumerate(trades, start=2):
    r[9] = f"=H{i}*I{i}"
sheet(tb, ["Date", "Time", "Trade ID", "Broker Order ID", "Exchange Order ID", "Symbol", "Side", "Qty", "Price",
           "Value", "Order Limit", "User Tag"],
      trades, [11, 10, 10, 15, 20, 28, 7, 8, 10, 14, 11, 52], money=(9, 10, 11))

# Orders: one row per broker order, its final state.
ow = wb.create_sheet("Orders")
orows = []
for oid, o in orders.items():
    avg = o["value"] / o["filled"] if o["filled"] else None
    orows.append([o["date"], o["first"], o.get("last", ""), oid, o["symbol"], o["side"], o["type"], o["qty"],
                  o["limit"], o["status"], o["filled"], avg, o["reason"], o["tag"]])
sheet(ow, ["Date", "First Seen", "Last Update", "Broker Order ID", "Symbol", "Side", "Type", "Qty", "Limit Price",
           "Final Status", "Filled Qty", "Avg Fill Price", "Reject / Remark", "User Tag"],
      orows, [11, 10, 11, 15, 28, 7, 8, 8, 11, 14, 10, 13, 70, 52], money=(9, 12))

ol = wb.create_sheet("Order Log")
sheet(ol, ["Date", "Time", "Event", "Broker Order ID", "Exchange Order ID", "Symbol", "Side", "Qty", "Price", "Type",
           "Status", "Pending Qty", "Filled Today", "Reason", "User Tag"],
      order_log, [11, 10, 22, 15, 20, 28, 7, 8, 10, 8, 14, 11, 11, 70, 52], money=(9,))

ev = wb.create_sheet("Platform Events")
sheet(ev, ["Date", "Time", "Event (execution gateway log)"], events, [11, 10, 180])

# Summary: totals as formulas over the Trade Book, and net per symbol.
sm = wb.create_sheet("Summary", 0)
n = len(trades) + 1
rng = lambda col: f"'Trade Book'!${col}$2:${col}${max(n, 2)}"
sm.append(["Orders and trades", day])
sm.append(["Source", f"broker order / trade pushes (Iris) in {os.path.join(log_dir, '6_reconciler.log')}"])
sm.append([])
sm.append(["Broker orders", len(orders)])
sm.append(["Trades (fills)", f"=COUNTA({rng('C')})"])
sm.append(["Sold qty", f'=SUMIFS({rng("H")},{rng("G")},"SELL")'])
sm.append(["Bought qty", f'=SUMIFS({rng("H")},{rng("G")},"BUY")'])
sm.append(["Sold value", f'=SUMIFS({rng("J")},{rng("G")},"SELL")'])
sm.append(["Bought value", f'=SUMIFS({rng("J")},{rng("G")},"BUY")'])
sm.append(["Net cash (sold - bought), gross of charges", "=B8-B9"])
sm.append([])
sm.append(["Symbol", "Sold Qty", "Bought Qty", "Net Qty (bought - sold)", "Sold Value", "Bought Value",
           "Net Cash (sold - bought)"])
hdr = sm.max_row
for s in sorted({t[5] for t in trades}):
    r = sm.max_row + 1
    sm.append([s,
               f'=SUMIFS({rng("H")},{rng("F")},$A{r},{rng("G")},"SELL")',
               f'=SUMIFS({rng("H")},{rng("F")},$A{r},{rng("G")},"BUY")',
               f"=C{r}-B{r}",
               f'=SUMIFS({rng("J")},{rng("F")},$A{r},{rng("G")},"SELL")',
               f'=SUMIFS({rng("J")},{rng("F")},$A{r},{rng("G")},"BUY")',
               f"=E{r}-F{r}"])
for row in sm.iter_rows():
    for c in row:
        c.font = FONT
        if c.row >= 6 and c.column >= 2 and c.row != hdr:
            c.number_format = "#,##0.00"
for c in sm[hdr]:
    c.font, c.fill = BOLD, HEAD
sm["A1"].font = Font(name="Arial", size=12, bold=True)
for i, w in enumerate([44, 16, 14, 24, 16, 16, 26], 1):
    sm.column_dimensions[get_column_letter(i)].width = w

wb.save(out)
print(out)
print(f"orders={len(orders)} trades={len(trades)} order_events={len(order_log)} platform_events={len(events)}")
